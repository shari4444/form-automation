"""
Google Form Automation Script (Playwright)
Target Form: https://docs.google.com/forms/d/e/1FAIpQLScIxD14Pyc183PbDLHcS9N-nbQRRYSncqY0TP3g7qSsfWgnKQ/viewform
Data Source: Google Cloud Learn with Communities - Participant Enrolment Form (Responses).xlsx
"""

import os
import sys
import json
import time
import argparse
import logging
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.path.join(BASE_DIR, "config.json")

# Load local configuration
with open(CONFIG_FILE, "r", encoding="utf-8") as f:
    CONFIG = json.load(f)

FORM_URL = CONFIG["form_url"]
EXCEL_FILE = os.path.join(BASE_DIR, CONFIG["excel_file"])
LOG_DIR = os.path.join(BASE_DIR, "logs")
SCREENSHOT_DIR = os.path.join(LOG_DIR, "screenshots")
USER_DATA_DIR = os.path.join(BASE_DIR, "browser_session")

os.makedirs(LOG_DIR, exist_ok=True)
os.makedirs(SCREENSHOT_DIR, exist_ok=True)
os.makedirs(USER_DATA_DIR, exist_ok=True)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(os.path.join(LOG_DIR, "automation.log"), encoding="utf-8"),
        logging.StreamHandler(sys.stdout)
    ]
)
logger = logging.getLogger("form_automation")


def load_participants(excel_path=EXCEL_FILE, only_unfilled=True, batch_tag=None):
    """
    Reads participants from the Excel response sheet.
    """
    wb = load_workbook(excel_path, data_only=True)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=False))
    
    participants = []
    current_batch = None
    
    for row_idx, row in enumerate(rows[1:], start=2):
        col0_val = row[0].value
        if col0_val is not None and str(col0_val).strip():
            current_batch = str(col0_val).strip()
            
        tick_val = row[1].value  # Column B
        
        # Skip already filled entries
        if only_unfilled and tick_val is True:
            continue
            
        # Optional filter by batch
        if batch_tag and current_batch and batch_tag.lower() not in current_batch.lower():
            continue

        full_name = row[9].value
        email = row[3].value or row[4].value
        skills_email = row[13].value or email
        profile_url = row[14].value
        gender = row[10].value
        
        if not full_name or not email:
            continue

        participants.append({
            "row_number": row_idx,
            "batch": current_batch,
            "name": str(full_name).strip(),
            "email": str(email).strip(),
            "skills_email": str(skills_email).strip(),
            "gender": str(gender).strip() if gender else "Male",
            "profile_url": str(profile_url).strip() if profile_url else ""
        })

    logger.info(f"Loaded {len(participants)} pending participant(s) from Excel.")
    return participants


def click_next(page):
    """Clicks Next button and waits for transition."""
    btn = page.get_by_role("button", name="Next")
    btn.click()
    page.wait_for_timeout(1000)


def fill_participant(page, p, organiser_type, referral_code, learner_category="Student"):
    """
    Steps through all 7 pages of the Google Form and fills participant data.
    """
    page.goto(FORM_URL, wait_until="networkidle", timeout=60000)

    # Check if signed in
    if page.get_by_text("Sign in to continue").count() > 0 or page.get_by_role("button", name="Sign in").count() > 0:
        logger.warning("Google Form requires Google Sign-in! Please log into Google in the open browser window.")
        page.wait_for_selector('text=/Next/i', timeout=120000)

    # ---------------- PAGE 1: Email ----------------
    logger.info(f"Filling Page 1: Email -> {p['email']}")
    email_field = page.locator('input[type="text"], input[type="email"]').filter(has_not=page.locator('[readonly]')).first
    if email_field.count() > 0 and email_field.is_editable():
        email_field.fill(p["email"])
    click_next(page)

    # ---------------- PAGE 2: Terms and Conditions ----------------
    logger.info("Filling Page 2: Terms and conditions")
    page.get_by_text("Yes, I accept the terms").click()
    click_next(page)

    # ---------------- PAGE 3: Participant Details ----------------
    logger.info(f"Filling Page 3: Name={p['name']}, Gender={p['gender']}, Category={learner_category}")
    # Full Name input
    page.locator('input[type="text"]').first.fill(p["name"])
    
    # Gender Dropdown (Type 3)
    gender_label = "Female" if "female" in p["gender"].lower() else "Male"
    dropdowns = page.locator('div[role="listbox"]').all()
    if len(dropdowns) >= 1:
        dropdowns[0].click()
        page.wait_for_timeout(300)
        page.get_by_role("option", name=gender_label, exact=True).click()
        page.wait_for_timeout(400)
    else:
        page.get_by_text(gender_label, exact=True).click()
    
    # Learner Category Dropdown (Type 3)
    if len(dropdowns) >= 2:
        dropdowns[1].click()
        page.wait_for_timeout(300)
        page.get_by_role("option", name=learner_category, exact=False).click()
        page.wait_for_timeout(400)
    else:
        page.get_by_text(learner_category, exact=False).click()
        
    click_next(page)

    # ---------------- PAGE 4: Choose Organiser ----------------
    logger.info(f"Filling Page 4: Organiser={organiser_type}, Code={referral_code}")
    # Select Your Organiser Name is also a dropdown (Type 3)
    org_dropdown = page.locator('div[role="listbox"]')
    if org_dropdown.count() > 0:
        org_dropdown.first.click()
        page.wait_for_timeout(300)
        page.get_by_role("option", name=organiser_type, exact=True).click()
        page.wait_for_timeout(400)
    else:
        page.get_by_text(organiser_type, exact=True).click()

    # Referral code
    page.locator('input[type="text"]').first.fill(referral_code)
    
    # Working internet and laptop
    page.get_by_text("Yes", exact=True).click()
    click_next(page)

    # ---------------- PAGE 5: Google Skills Account Verification ----------------
    logger.info(f"Filling Page 5: Skills Email={p['skills_email']}")
    page.get_by_text("I verify").click()
    page.locator('input[type="text"]').first.fill(p["skills_email"])
    click_next(page)

    # ---------------- PAGE 6: Public Profile URL ----------------
    logger.info(f"Filling Page 6: Profile URL={p['profile_url']}")
    textarea = page.locator('textarea')
    if textarea.count() > 0:
        textarea.fill(p["profile_url"])
    else:
        page.locator('input[type="text"]').first.fill(p["profile_url"])
    click_next(page)

    # ---------------- PAGE 7: Acknowledge & Submit ----------------
    logger.info("Filling Page 7: Acknowledge & Submit")
    page.get_by_text("I Acknowledge").click()
    page.get_by_text("Yes, I Agree").click()

    # Submit button
    submit_btn = page.get_by_role("button", name="Submit")
    submit_btn.click()

    # Verify confirmation message on success screen
    page.wait_for_selector('text=/Thank you|response has been recorded/i', timeout=15000)
    logger.info(f"SUCCESS: Form submitted for {p['name']}!")


def mark_row_as_filled_in_excel(excel_path, row_number):
    try:
        wb = load_workbook(excel_path)
        sheet = wb.active
        sheet.cell(row=row_number, column=2, value=True)
        wb.save(excel_path)
        logger.info(f"Excel updated: row {row_number} (Column B = True).")
    except Exception as e:
        logger.error(f"Failed to update Excel row {row_number}: {e}")


def main():
    parser = argparse.ArgumentParser(description="Submit participant entries from Excel to Google Form.")
    parser.add_argument("--limit", type=int, default=1, help="Number of participants to submit (default: 1).")
    parser.add_argument("--batch", type=str, default=None, help="Filter by batch tag (e.g., 'darshan', 'p2', etc.).")
    parser.add_argument("--referral-code", type=str, default=CONFIG["referral_code"], help="Organizer Referral Code.")
    parser.add_argument("--organiser", type=str, default=CONFIG["organizer_type"])
    parser.add_argument("--category", type=str, default="Student")
    parser.add_argument("--login-only", action="store_true", help="Launch browser for you to sign into Google once, then exit.")
    args = parser.parse_args()

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=USER_DATA_DIR,
            headless=False,
            viewport={"width": 1280, "height": 900},
            args=["--disable-blink-features=AutomationControlled"]
        )
        page = context.new_page()

        if args.login_only:
            logger.info("Opening Google Form in browser for manual Google Sign-in...")
            page.goto(FORM_URL)
            logger.info("Please sign into your Google Account in the browser window.")
            logger.info("Once signed in and the form is visible, press Enter in the terminal to finish setup.")
            input("Press Enter here once you have signed into Google: ")
            context.close()
            logger.info("Session saved to 'browser_session/'! You can now run form submissions.")
            return

        participants = load_participants(only_unfilled=True, batch_tag=args.batch)
        to_process = participants[:args.limit]

        success_count = 0
        fail_count = 0

        for i, participant in enumerate(to_process, 1):
            logger.info(f"[{i}/{len(to_process)}] Processing Row {participant['row_number']}: {participant['name']} ({participant['email']})")
            try:
                fill_participant(page, participant, args.organiser, args.referral_code, args.category)
                mark_row_as_filled_in_excel(EXCEL_FILE, participant["row_number"])
                success_count += 1
                time.sleep(2)
            except Exception as e:
                fail_count += 1
                logger.error(f"Error submitting row {participant['row_number']} ({participant['name']}): {e}")
                ss_path = os.path.join(SCREENSHOT_DIR, f"error_row_{participant['row_number']}.png")
                try:
                    page.screenshot(path=ss_path)
                    logger.info(f"Saved screenshot: {ss_path}")
                except Exception:
                    pass

        context.close()
        logger.info(f"Completed run: {success_count} succeeded, {fail_count} failed.")


if __name__ == "__main__":
    main()
